"""Real Office conversion with fictitious documents; nothing is uploaded."""
import importlib.util
import io
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

ROOT = Path(__file__).parents[1]
spec = importlib.util.spec_from_file_location('office', ROOT / 'local/office.py')
office = importlib.util.module_from_spec(spec)
spec.loader.exec_module(office)

def docx():
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Servicio de prueba 210 EUR</w:t></w:r></w:p></w:body></w:document>')
    return out.getvalue()

def xlsx():
    workbook = Workbook()
    workbook.active.append(['Servicio', 'Importe'])
    workbook.active.append(['Prueba', 210])
    workbook.active.print_options.horizontalCentered = True
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()

def pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = 'Servicio de prueba 210 EUR'
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()

class OfficeTests(unittest.TestCase):
    def test_modern_formats_validate_and_macro_payloads_fail(self):
        for extension, create in [('docx', docx), ('xlsx', xlsx), ('pptx', pptx)]:
            office.validate_office(create(), extension)
        bad = io.BytesIO(docx())
        with zipfile.ZipFile(bad, 'a') as z:
            z.writestr('word/vbaProject.bin', b'fake')
        with self.assertRaises(office.OfficeError):
            office.validate_office(bad.getvalue(), 'docx')

    @unittest.skipUnless(office.find_office(), 'LibreOffice no instalado')
    def test_real_docx_xlsx_pptx_with_both_desktop_adapters(self):
        for extension, create in [('docx', docx), ('xlsx', xlsx), ('pptx', pptx)]:
            data = create()
            with self.subTest(extension=extension, adapter='python'):
                result = office.convert_office(data, extension)
                reader = PdfReader(io.BytesIO(result))
                self.assertEqual(len(reader.pages), 1)
                self.assertIn('210', reader.pages[0].extract_text())
            with self.subTest(extension=extension, adapter='node'), tempfile.TemporaryDirectory() as folder:
                source = Path(folder, 'test.' + extension)
                source.write_bytes(data)
                code = "import {readFile,writeFile} from 'node:fs/promises';import {convertOffice} from './local/office.mjs';const result=await convertOffice(await readFile(process.argv[1]),process.argv[2]);await writeFile(process.argv[3],result);"
                dest = Path(folder, 'out.pdf')
                process = subprocess.run(['node', '--input-type=module', '-e', code, str(source), extension, str(dest)], cwd=ROOT, capture_output=True, timeout=110)
                self.assertEqual(process.returncode, 0, process.stderr.decode())
                reader = PdfReader(dest)
                self.assertEqual(len(reader.pages), 1)
                self.assertIn('210', reader.pages[0].extract_text())

    def test_missing_engine_and_busy_are_explained(self):
        with patch.object(office, 'find_office', return_value=None):
            with self.assertRaises(office.OfficeError) as error:
                office.convert_office(docx(), 'docx')
            self.assertEqual(error.exception.code, 'OFFICE_MISSING')
        with patch.object(office, 'find_office', return_value='/fake/soffice'):
            office.LOCK.acquire()
            try:
                with self.assertRaises(office.OfficeError) as error:
                    office.convert_office(docx(), 'docx')
                self.assertEqual(error.exception.code, 'OFFICE_BUSY')
            finally:
                office.LOCK.release()
